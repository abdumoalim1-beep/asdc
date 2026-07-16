import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from asdc import agent, server


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "registry", server.WorkspaceRegistry(tmp_path / "workspaces"))
    monkeypatch.setattr(server, "_model_call", None)
    return TestClient(server.app)


def _new_workspace(api, name="شركة تجريبية"):
    res = api.post("/api/workspaces", json={"name": name})
    assert res.status_code == 200
    return res.json()["id"]


def _tool_call(call_id, name, arguments):
    return {"id": call_id, "type": "function", "function": {"name": name, "arguments": json.dumps(arguments, ensure_ascii=False)}}


def _set_fake_model(monkeypatch, fn):
    monkeypatch.setattr(server, "_model_call", fn)


def test_create_workspace_starts_empty(api):
    ws_id = _new_workspace(api)
    res = api.get(f"/api/workspaces/{ws_id}/state")
    assert res.status_code == 200
    data = res.json()
    assert data["is_new"] is True
    assert data["stats"] == {}
    assert data["files"] == []
    assert data["workspace_name"] == "شركة تجريبية"


def test_unknown_workspace_returns_404(api):
    res = api.get("/api/workspaces/does-not-exist/state")
    assert res.status_code == 404


def test_two_workspaces_do_not_share_data(api, monkeypatch):
    ws_a = _new_workspace(api, "أ")
    ws_b = _new_workspace(api, "ب")

    def fake_model(messages):
        tool_msgs = [m for m in messages if m.get("role") == "tool"]
        if not tool_msgs:
            return {"role": "assistant", "content": "", "tool_calls": [_tool_call("c1", "add_entity", {"entity_type": "عميل", "name": "عميل أ", "fields": {}})]}
        return {"role": "assistant", "content": "تمام، أضفته."}

    _set_fake_model(monkeypatch, fake_model)
    api.post(f"/api/workspaces/{ws_a}/chat", json={"message": "عندي عميل اسمه عميل أ"})

    state_a = api.get(f"/api/workspaces/{ws_a}/state").json()
    state_b = api.get(f"/api/workspaces/{ws_b}/state").json()
    assert state_a["stats"] == {"عميل": 1}
    assert state_b["stats"] == {}


def test_chat_plain_reply_no_tool(api, monkeypatch):
    ws_id = _new_workspace(api)

    def fake_model(messages):
        return {"role": "assistant", "content": "أهلًا فيك! عرّفني بشركتك عشان أبدأ."}

    _set_fake_model(monkeypatch, fake_model)
    res = api.post(f"/api/workspaces/{ws_id}/chat", json={"message": "مرحبا"})
    assert res.status_code == 200
    data = res.json()
    assert data["reply"] == "أهلًا فيك! عرّفني بشركتك عشان أبدأ."
    assert data["actions"] == []


def test_chat_executes_tool_and_persists_conversation(api, monkeypatch):
    ws_id = _new_workspace(api)
    calls = []

    def fake_model(messages):
        calls.append(messages)
        tool_msgs = [m for m in messages if m.get("role") == "tool"]
        if not tool_msgs:
            return {"role": "assistant", "content": "", "tool_calls": [_tool_call("c1", "add_entity", {"entity_type": "منتج", "name": "باقة تجريبية", "fields": {"السعر": "100"}})]}
        return {"role": "assistant", "content": "أضفت باقة تجريبية بسعر 100."}

    _set_fake_model(monkeypatch, fake_model)
    res = api.post(f"/api/workspaces/{ws_id}/chat", json={"message": "أضف منتج اسمه باقة تجريبية بسعر 100"})
    data = res.json()
    assert data["actions"][0]["tool"] == "add_entity"
    assert data["actions"][0]["ok"] is True

    state = api.get(f"/api/workspaces/{ws_id}/state").json()
    assert state["stats"] == {"منتج": 1}
    assert state["is_new"] is False

    # second turn should include the first turn's history for the model
    api.post(f"/api/workspaces/{ws_id}/chat", json={"message": "شكرا"})
    assert len(calls) >= 2
    roles_in_second_call = [m["role"] for m in calls[-1]]
    assert "assistant" in roles_in_second_call


def test_chat_tool_failure_is_reported_not_crashed(api, monkeypatch):
    ws_id = _new_workspace(api)

    def fake_model(messages):
        tool_msgs = [m for m in messages if m.get("role") == "tool"]
        if not tool_msgs:
            return {"role": "assistant", "content": "", "tool_calls": [_tool_call("c1", "update_field", {"entity_id": "no-such-id", "fields": {"السعر": "1"}})]}
        payload = json.loads(tool_msgs[0]["content"])
        assert payload["ok"] is False
        return {"role": "assistant", "content": "ما لقيت هذا الكيان."}

    _set_fake_model(monkeypatch, fake_model)
    res = api.post(f"/api/workspaces/{ws_id}/chat", json={"message": "حدث كيان غير موجود"})
    data = res.json()
    assert data["actions"][0]["ok"] is False
    assert data["actions"][0]["error"]
    assert "ما لقيت" in data["reply"]


def test_upload_template_with_explicit_placeholders_needs_no_llm(api, tmp_path):
    ws_id = _new_workspace(api)
    docx_path = tmp_path / "raw.docx"
    doc = Document()
    doc.add_paragraph("عقد للعميل: {{الاسم}}")
    doc.save(str(docx_path))

    with open(docx_path, "rb") as fh:
        res = api.post(
            f"/api/workspaces/{ws_id}/templates/upload",
            files={"file": ("raw.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"name": "عقد اختباري", "applies_to": "عميل"},
        )
    assert res.status_code == 200
    data = res.json()
    assert data["fields"] == ["الاسم"]

    _entities, templates = server.registry.load_stores(ws_id)
    assert templates.get("عقد اختباري") is not None


def test_upload_template_is_scoped_to_its_workspace(api, tmp_path):
    ws_a = _new_workspace(api, "أ")
    ws_b = _new_workspace(api, "ب")
    docx_path = tmp_path / "raw.docx"
    doc = Document()
    doc.add_paragraph("عقد للعميل: {{الاسم}}")
    doc.save(str(docx_path))

    with open(docx_path, "rb") as fh:
        api.post(
            f"/api/workspaces/{ws_a}/templates/upload",
            files={"file": ("raw.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"name": "عقد أ", "applies_to": "عميل"},
        )

    _entities_a, templates_a = server.registry.load_stores(ws_a)
    _entities_b, templates_b = server.registry.load_stores(ws_b)
    assert templates_a.get("عقد أ") is not None
    assert templates_b.get("عقد أ") is None
