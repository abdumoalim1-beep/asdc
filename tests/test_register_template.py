from pathlib import Path

import pytest
from docx import Document

from asdc.register_template import TemplateRegistrationError, register_docx_template
from asdc.templates import TemplateStore


def _make_docx(path: Path, text: str) -> Path:
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def test_register_with_explicit_placeholders(tmp_path):
    src = _make_docx(tmp_path / "src.docx", "عقد للعميل: {{الاسم}}")
    store = TemplateStore([])
    template = register_docx_template(src, "عقد اختباري", "عميل", tmp_path / "stored", store)
    assert template.fields == ["الاسم"]
    assert Path(template.source_path).exists()
    assert store.get("عقد اختباري") is template


def test_register_without_placeholders_requires_completion_fn(tmp_path):
    src = _make_docx(tmp_path / "src.docx", "الطرف الثاني: شركة الرياض للتسويق")
    store = TemplateStore([])
    with pytest.raises(TemplateRegistrationError):
        register_docx_template(src, "عقد اختباري", "عميل", tmp_path / "stored", store)


def test_register_without_placeholders_uses_llm_suggestions(tmp_path):
    src = _make_docx(tmp_path / "src.docx", "الطرف الثاني: شركة الرياض للتسويق")
    store = TemplateStore([])

    def fake_completion(system_prompt, payload):
        return '[{"example_text": "شركة الرياض للتسويق", "field_name": "الاسم"}]'

    template = register_docx_template(
        src, "عقد اختباري", "عميل", tmp_path / "stored", store, completion_fn=fake_completion
    )
    assert template.fields == ["الاسم"]
    assert "{{الاسم}}" in Document(template.source_path).paragraphs[0].text
