from pathlib import Path

import pytest
from docx import Document

from asdc import docx_template
from asdc.llm_client import MalformedModelResponse


def _make_docx(path: Path, paragraphs: list[list[str]]) -> Path:
    """paragraphs: list of paragraphs, each a list of run texts to add separately."""
    doc = Document()
    for runs in paragraphs:
        p = doc.add_paragraph()
        for text in runs:
            p.add_run(text)
    doc.save(str(path))
    return path


def test_extract_placeholders_single_run(tmp_path):
    path = _make_docx(tmp_path / "t.docx", [["عقد للعميل: {{الاسم}}، البريد: {{البريد الإلكتروني}}"]])
    assert docx_template.extract_placeholders(path) == ["الاسم", "البريد الإلكتروني"]


def test_extract_placeholders_split_across_runs(tmp_path):
    path = _make_docx(tmp_path / "t.docx", [["عقد للعميل: {{", "الاسم", "}}"]])
    assert docx_template.extract_placeholders(path) == ["الاسم"]


def test_build_rendered_docx_fills_values(tmp_path):
    path = _make_docx(tmp_path / "t.docx", [["الاسم: {{الاسم}}"], ["السعر: {{السعر}} ريال"]])
    doc, missing = docx_template.build_rendered_docx(path, {"الاسم": "شركة الرياض", "السعر": "500"})
    assert missing == []
    texts = [p.text for p in doc.paragraphs]
    assert "الاسم: شركة الرياض" in texts
    assert "السعر: 500 ريال" in texts


def test_build_rendered_docx_reports_missing_without_inventing(tmp_path):
    path = _make_docx(tmp_path / "t.docx", [["الاسم: {{الاسم}}، الهاتف: {{رقم الجوال}}"]])
    doc, missing = docx_template.build_rendered_docx(path, {"الاسم": "شركة الرياض"})
    assert missing == ["رقم الجوال"]
    # unresolved placeholder text is left untouched, not filled with a guess
    assert "{{رقم الجوال}}" in doc.paragraphs[0].text


def test_extract_full_text(tmp_path):
    path = _make_docx(tmp_path / "t.docx", [["سطر أول"], ["سطر ثاني"]])
    text = docx_template.extract_full_text(path)
    assert "سطر أول" in text and "سطر ثاني" in text


def test_suggest_placeholders_extractive_check_drops_hallucinated_example(tmp_path):
    template_text = "الطرف الثاني: شركة الرياض للتسويق"

    def fake_completion(system_prompt, payload):
        return (
            '[{"example_text": "شركة الرياض للتسويق", "field_name": "الاسم"}, '
            '{"example_text": "نص غير موجود إطلاقًا", "field_name": "وهمي"}]'
        )

    suggestions = docx_template.suggest_placeholders_with_llm(template_text, fake_completion)
    assert suggestions == [{"example_text": "شركة الرياض للتسويق", "field_name": "الاسم"}]


def test_suggest_placeholders_rejects_non_array_reply():
    def fake_completion(system_prompt, payload):
        return '{"not": "an array"}'

    with pytest.raises(MalformedModelResponse):
        docx_template.suggest_placeholders_with_llm("نص", fake_completion)


def test_apply_suggestions_to_docx_inserts_placeholders(tmp_path):
    src = _make_docx(tmp_path / "src.docx", [["الطرف الثاني: شركة الرياض للتسويق"]])
    out = tmp_path / "out.docx"
    inserted = docx_template.apply_suggestions_to_docx(
        src, [{"example_text": "شركة الرياض للتسويق", "field_name": "الاسم"}], out
    )
    assert inserted == ["الاسم"]
    assert docx_template.extract_placeholders(out) == ["الاسم"]
