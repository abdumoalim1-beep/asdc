from pathlib import Path

import pytest
from docx import Document

from asdc.executor import ExecutionRefused, execute
from asdc.schema import NLUResponse
from asdc.templates import Template


def _make_docx_template(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("عقد تقديم خدمات")
    doc.add_paragraph("الطرف الثاني: {{الاسم}}")
    doc.add_paragraph("البريد الإلكتروني: {{البريد الإلكتروني}}")
    doc.save(str(path))
    return path


def _response(**overrides):
    base = dict(
        intent="generate_document",
        confidence="high",
        entity_type="عميل",
        entity_reference=None,
        matched_entity_id=None,
        fields_to_update={},
        template_requested="عقد Word حقيقي",
        clarification_needed=None,
        human_message="جاري التجهيز.",
    )
    base.update(overrides)
    return NLUResponse.model_validate(base)


def test_generate_document_fills_and_saves_real_docx(entities, templates, documents_dir, tmp_path):
    docx_path = _make_docx_template(tmp_path / "template.docx")
    templates.add(Template(name="عقد Word حقيقي", applies_to="عميل", fields=["الاسم", "البريد الإلكتروني"],
                            source_path=str(docx_path)))

    response = _response(matched_entity_id="client_003")
    result = execute(response, entities, templates, documents_dir)

    output_path = documents_dir / f"{result.detail['document_id']}.docx"
    assert output_path.exists()
    generated_text = "\n".join(p.text for p in Document(str(output_path)).paragraphs)
    assert "شركة الرياض للتسويق" in generated_text
    assert "info@riyadh-marketing.example.com" in generated_text


def test_generate_document_docx_refuses_when_field_missing(entities, templates, documents_dir, tmp_path):
    docx_path = _make_docx_template(tmp_path / "template.docx")
    templates.add(Template(name="عقد Word حقيقي", applies_to="مؤثر", fields=["الاسم", "البريد الإلكتروني"],
                            source_path=str(docx_path)))

    response = _response(matched_entity_id="infl_007", entity_type="مؤثر")
    with pytest.raises(ExecutionRefused):
        execute(response, entities, templates, documents_dir)
    assert entities.get("infl_007").documents == []


def test_generate_document_docx_new_version_for_signed_entity(entities, templates, documents_dir, tmp_path):
    docx_path = _make_docx_template(tmp_path / "template.docx")
    templates.add(Template(name="عقد Word حقيقي", applies_to="عميل", fields=["الاسم", "البريد الإلكتروني"],
                            source_path=str(docx_path)))

    entity = entities.get("client_009")
    original_signed = dict(entity.documents[0])

    response = _response(matched_entity_id="client_009")
    result = execute(response, entities, templates, documents_dir)

    assert result.detail["is_new_version"] is True
    assert result.detail["superseded_status"] == "موقّع"
    assert entity.documents[0] == original_signed
    assert (documents_dir / f"{result.detail['document_id']}.docx").exists()
